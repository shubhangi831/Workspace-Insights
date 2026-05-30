# Copyright (c) 2026, sk and contributors
# For license information, please see license.txt

# import frappe
# from frappe.model.document import Document


# class SaleInvoice(Document):
# 	pass


# import frappe
# import csv
# import io
# import base64
# import unicodedata
# from collections import defaultdict
# from frappe.model.document import Document
# from frappe.utils import flt, nowdate


# class SaleInvoice(Document):

#     def before_save(self):
#         """Server-side GST recalculation — safety net if JS missed anything."""
#         self.calculate_gst()

#     def calculate_gst(self):
#         subtotal = sum(flt(row.amount) for row in (self.domain_details or []))
#         gst_rate = flt(self.gst_rate) or 18

#         self.subtotal       = flt(subtotal, 2)
#         self.gst_amount     = flt(subtotal * gst_rate / 100, 2)
#         self.grand_total    = flt(self.subtotal + self.gst_amount, 2)
#         self.invoice_amount = self.grand_total


# # ── xlsx → CSV ────────────────────────────────────────────────────

# def xlsx_to_csv(b64_content):
#     try:
#         import openpyxl
#     except ImportError:
#         frappe.throw("Run: pip install openpyxl --break-system-packages")
#     raw    = base64.b64decode(b64_content)
#     wb     = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
#     ws     = wb.active
#     output = io.StringIO()
#     writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
#     for row in ws.iter_rows(values_only=True):
#         writer.writerow(['' if c is None else str(c) for c in row])
#     return output.getvalue()


# # ── NFKC normalization ────────────────────────────────────────────

# def norm(s):
#     s = unicodedata.normalize('NFKC', s or '')
#     return ' '.join(s.split()).lower()


# # ── Main import method ────────────────────────────────────────────

# @frappe.whitelist()
# def import_sale_invoice_csv(csv_content='', file_type='csv'):
#     """
#     Import Sale Invoice from a flat CSV/XLSX file.

#     Required columns: invoice_number, invoice_date, domain,
#                       subscription, quantity, amount
#     Optional columns: description, start_date, end_date,
#                       order_name, po_number, licenses, bill_to

#     Rows grouped by invoice_number → one Sale Invoice per group.
#     GST is calculated automatically via before_save.
#     """
#     try:
#         if file_type == 'xlsx':
#             try:
#                 csv_content = xlsx_to_csv(csv_content)
#             except Exception as e:
#                 return {'success': False, 'error': f"Cannot read Excel: {e}"}

#         if csv_content.startswith('\ufeff'):
#             csv_content = csv_content[1:]

#         lines = csv_content.replace('\r', '').split('\n')
#         lines = [l for l in lines if l.strip()]

#         if len(lines) < 2:
#             return {'success': False, 'error': 'File is empty or has no data rows.'}

#         # Parse header row
#         header_row = [c.strip().lower().replace(' ', '_')
#                       for c in next(csv.reader(io.StringIO(lines[0])))]

#         required = {'invoice_number', 'invoice_date', 'domain',
#                     'subscription', 'quantity', 'amount'}
#         missing_cols = required - set(header_row)
#         if missing_cols:
#             return {'success': False,
#                     'error': f"Missing required columns: {', '.join(sorted(missing_cols))}"}

#         def col(row_dict, name, default=''):
#             v = row_dict.get(name, default)
#             return unicodedata.normalize('NFKC', str(v) if v else '').strip()

#         # Build NFKC-normalized subscription plan lookup
#         all_plans   = frappe.get_all('Subscription Plan', fields=['name', 'subscription'])
#         plan_lookup = {}
#         for p in all_plans:
#             if p.subscription: plan_lookup[norm(p.subscription)] = p.name
#             if p.name:         plan_lookup[norm(p.name)]         = p.name

#         # Parse data rows
#         raw_rows = []
#         for line in lines[1:]:
#             if not line.strip(): continue
#             try:
#                 values   = next(csv.reader(io.StringIO(line)))
#                 row_dict = {header_row[i]: values[i].strip()
#                             if i < len(values) else ''
#                             for i in range(len(header_row))}
#                 raw_rows.append(row_dict)
#             except Exception:
#                 continue

#         # Group by invoice_number
#         invoice_groups = defaultdict(list)
#         for row in raw_rows:
#             inv_num = col(row, 'invoice_number')
#             if inv_num:
#                 invoice_groups[inv_num].append(row)

#         if not invoice_groups:
#             return {'success': False, 'error': 'No valid invoice rows found.'}

#         created    = []
#         duplicates = []
#         failed     = []
#         all_miss_d = []
#         all_miss_s = []

#         for inv_num, rows in invoice_groups.items():
#             first = rows[0]

#             # Duplicate check
#             existing = frappe.db.get_value('Sale Invoice', {'invoice_number': inv_num}, 'name')
#             if existing:
#                 duplicates.append({'invoice_number': inv_num, 'existing_doc': existing})
#                 continue

#             missing_domains       = []
#             missing_subscriptions = []
#             valid_rows            = []

#             for row in rows:
#                 domain       = col(row, 'domain')
#                 subscription = col(row, 'subscription')
#                 if not domain or not subscription:
#                     continue

#                 # Domain check
#                 if not frappe.db.exists('Domains', domain):
#                     if domain not in missing_domains:
#                         missing_domains.append(domain)
#                     continue

#                 # Subscription check — store plan NAME for Link validation
#                 plan_name = plan_lookup.get(norm(subscription))
#                 if not plan_name:
#                     if subscription not in missing_subscriptions:
#                         missing_subscriptions.append(subscription)
#                     continue

#                 def parse_date(raw):
#                     if not raw: return None
#                     import re
#                     raw = raw.strip()
#                     months = {m: str(i+1).zfill(2) for i, m in enumerate(
#                         ['jan','feb','mar','apr','may','jun','jul','aug',
#                          'sep','oct','nov','dec'])}
#                     if re.match(r'\d{4}-\d{2}-\d{2}', raw): return raw
#                     for sep in ['-', '/']:
#                         parts = raw.split(sep)
#                         if len(parts) == 3:
#                             if len(parts[0]) == 4:
#                                 return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
#                             yr = ('20'+parts[2]) if len(parts[2])==2 else parts[2]
#                             return f"{yr}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
#                     p = raw.split()
#                     if len(p) == 3:
#                         mon = months.get(p[1][:3].lower(), '01')
#                         yr  = ('20'+p[2]) if len(p[2])==2 else p[2]
#                         return f"{yr}-{mon}-{p[0].zfill(2)}"
#                     return None

#                 valid_rows.append({
#                     'doctype':      'Sale Invoice Items',
#                     'domain':       domain,
#                     'subscription': plan_name,
#                     'description':  col(row, 'description'),
#                     'order_name':   col(row, 'order_name'),
#                     'start_date':   parse_date(col(row, 'start_date')),
#                     'end_date':     parse_date(col(row, 'end_date')),
#                     'quantity':     int(col(row, 'quantity') or 0),
#                     'licenses':     int(col(row, 'licenses') or 0),
#                     'po_number':    col(row, 'po_number'),
#                     'amount':       float(col(row, 'amount').replace(',', '') or 0),
#                 })

#             if not valid_rows:
#                 failed.append({
#                     'invoice_number':        inv_num,
#                     'missing_domains':       missing_domains,
#                     'missing_subscriptions': missing_subscriptions,
#                     'error': 'No valid rows — all skipped due to missing domain/subscription.'
#                 })
#                 all_miss_d.extend(missing_domains)
#                 all_miss_s.extend(missing_subscriptions)
#                 continue

#             try:
#                 doc = frappe.get_doc({
#                     'doctype':        'Sale Invoice',
#                     'bill_to':        col(first, 'bill_to'),
#                     'invoice_number': inv_num,
#                     'invoice_date':   parse_date(col(first, 'invoice_date')),
#                     'currency':       'INR',
#                     'gst_rate':       18,
#                     'domain_details': valid_rows
#                 })
#                 doc.insert(ignore_permissions=True)
#                 frappe.db.commit()

#                 created.append({
#                     'doc':                   doc.name,
#                     'invoice_number':        inv_num,
#                     'total_rows':            len(valid_rows),
#                     'grand_total':           doc.grand_total,
#                     'missing_domains':       missing_domains,
#                     'missing_subscriptions': missing_subscriptions,
#                 })
#                 all_miss_d.extend(missing_domains)
#                 all_miss_s.extend(missing_subscriptions)

#             except Exception as e:
#                 frappe.log_error(frappe.get_traceback(), 'Sale Invoice Import Error')
#                 failed.append({'invoice_number': inv_num, 'error': str(e)})

#         return {
#             'success':               True,
#             'created':               created,
#             'duplicates':            duplicates,
#             'failed':                failed,
#             'missing_domains':       list(set(all_miss_d)),
#             'missing_subscriptions': list(set(all_miss_s)),
#         }

#     except Exception as e:
#         frappe.log_error(frappe.get_traceback(), 'Sale Invoice Import Error')
#         return {'success': False, 'error': str(e)}





# new import code

import frappe
import csv
import io
import base64
import unicodedata
from collections import defaultdict
from frappe.model.document import Document
from frappe.utils import flt, nowdate


class SaleInvoice(Document):

    def before_save(self):
        self.calculate_gst()

    def calculate_gst(self):
        subtotal = sum(flt(row.amount) for row in (self.domain_details or []))
        gst_rate = flt(self.gst_rate) or 18
        self.subtotal       = flt(subtotal, 2)
        self.gst_amount     = flt(subtotal * gst_rate / 100, 2)
        self.grand_total    = flt(self.subtotal + self.gst_amount, 2)
        self.invoice_amount = self.grand_total


# ── xlsx → CSV ────────────────────────────────────────────────────
def xlsx_to_csv(b64_content):
    try:
        import openpyxl
    except ImportError:
        frappe.throw("Run: pip install openpyxl --break-system-packages")
    raw    = base64.b64decode(b64_content)
    wb     = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws     = wb.active
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(['' if c is None else str(c) for c in row])
    return output.getvalue()


# ── NFKC normalization ────────────────────────────────────────────
def norm(s):
    s = unicodedata.normalize('NFKC', s or '')
    return ' '.join(s.split()).lower()


# ── Find matching PDF by invoice number ───────────────────────────
def find_matching_pdf(inv_num, pdf_files):
    if not inv_num or not pdf_files:
        return None
    for fname in pdf_files:
        if inv_num in fname:
            return fname
    return None


# ── Attach PDF to Sale Invoice ────────────────────────────────────
def attach_pdf_to_doc(doc_name, pdf_filename, pdf_base64):
    try:
        from frappe.utils.file_manager import save_file

        pdf_content = base64.b64decode(pdf_base64)

        file_doc = save_file(
            fname      = pdf_filename,
            content    = pdf_content,
            dt         = 'Sale Invoice',
            dn         = doc_name,
            df         = 'invoice_attachment',
            is_private = 1
        )

        frappe.db.set_value(
            'Sale Invoice', doc_name,
            'invoice_attachment', file_doc.file_url
        )
        frappe.db.commit()
        return True

    except Exception:
        frappe.log_error(frappe.get_traceback(), f'PDF Attach Error: {doc_name}')
        return False


# ── Main import method ────────────────────────────────────────────
@frappe.whitelist()
def import_sale_invoice_csv(csv_content='', file_type='csv', pdf_files=None):
    try:
        if isinstance(pdf_files, str):
            import json
            pdf_files = json.loads(pdf_files) if pdf_files else {}
        pdf_files = pdf_files or {}

        if file_type == 'xlsx':
            try:
                csv_content = xlsx_to_csv(csv_content)
            except Exception as e:
                return {'success': False, 'error': f"Cannot read Excel file: {e}"}

        if csv_content.startswith('\ufeff'):
            csv_content = csv_content[1:]

        lines = csv_content.replace('\r', '').split('\n')
        lines = [l for l in lines if l.strip()]

        if len(lines) < 2:
            return {'success': False, 'error': 'File is empty or has no data rows.'}

        header_row = [c.strip().lower().replace(' ', '_')
                      for c in next(csv.reader(io.StringIO(lines[0])))]

        # Skip marker rows (Required / Optional / Auto)
        data_lines = []
        for line in lines[1:]:
            if not line.strip():
                continue
            first_cell = next(csv.reader(io.StringIO(line)))[0].strip().lower()
            if first_cell in ('required', 'optional', 'auto'):
                continue
            data_lines.append(line)

        required = {'invoice_number', 'invoice_date', 'domain', 'subscription', 'quantity', 'amount'}
        missing_cols = required - set(header_row)
        if missing_cols:
            return {
                'success': False,
                'error':   f"Missing required columns: {', '.join(sorted(missing_cols))}"
            }

        def col(row_dict, name, default=''):
            v = row_dict.get(name, default)
            return unicodedata.normalize('NFKC', str(v) if v else '').strip()

        def parse_date(raw):
            if not raw: return None
            import re
            raw = raw.strip()
            if re.match(r'\d{4}-\d{2}-\d{2}', raw): return raw
            months = {m: str(i+1).zfill(2) for i, m in enumerate(
                ['jan','feb','mar','apr','may','jun','jul',
                 'aug','sep','oct','nov','dec'])}
            for sep in ['-', '/']:
                parts = raw.split(sep)
                if len(parts) == 3:
                    if len(parts[0]) == 4:
                        return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
                    yr = ('20'+parts[2]) if len(parts[2])==2 else parts[2]
                    return f"{yr}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
            p = raw.split()
            if len(p) == 3:
                mon = months.get(p[1][:3].lower(), '01')
                yr  = ('20'+p[2]) if len(p[2])==2 else p[2]
                return f"{yr}-{mon}-{p[0].zfill(2)}"
            return None

        # Subscription plan lookup — name, subscription, short_label all mein match
        all_plans   = frappe.get_all('Subscription Plan', fields=['name', 'subscription', 'short_label'])
        plan_lookup = {}
        for p in all_plans:
            if p.name:         plan_lookup[norm(p.name)]         = p.name
            if p.subscription: plan_lookup[norm(p.subscription)] = p.name
            if p.short_label:  plan_lookup[norm(p.short_label)]  = p.name

        # Parse data rows
        raw_rows = []
        for line in data_lines:
            try:
                values   = next(csv.reader(io.StringIO(line)))
                row_dict = {
                    header_row[i]: values[i].strip() if i < len(values) else ''
                    for i in range(len(header_row))
                }
                raw_rows.append(row_dict)
            except Exception:
                continue

        # Group by invoice_number
        invoice_groups = defaultdict(list)
        inv_meta       = {}
        current_inv    = None

        for row in raw_rows:
            inv_num = col(row, 'invoice_number')
            if inv_num:
                current_inv       = inv_num
                inv_meta[inv_num] = row
            if current_inv:
                invoice_groups[current_inv].append(row)

        if not invoice_groups:
            return {'success': False, 'error': 'No valid invoice rows found.'}

        created    = []
        duplicates = []
        failed     = []
        all_miss_d = []
        all_miss_s = []

        for inv_num, rows in invoice_groups.items():
            first = inv_meta.get(inv_num, rows[0])

            existing = frappe.db.get_value('Sale Invoice', {'invoice_number': inv_num}, 'name')
            if existing:
                duplicates.append({'invoice_number': inv_num, 'existing_doc': existing})
                continue

            missing_domains       = []
            missing_subscriptions = []
            valid_rows            = []

            for row in rows:
                domain       = col(row, 'domain')
                subscription = col(row, 'subscription')
                if not domain or not subscription:
                    continue

                if not frappe.db.exists('Domains', domain):
                    if domain not in missing_domains:
                        missing_domains.append(domain)
                    continue

                plan_name = plan_lookup.get(norm(subscription))
                if not plan_name:
                    if subscription not in missing_subscriptions:
                        missing_subscriptions.append(subscription)
                    continue

                qty_raw = col(row, 'quantity').replace(',', '').strip()
                amt_raw = col(row, 'amount').replace(',', '').strip()
                lic_raw = col(row, 'licenses').replace(',', '').strip()

                valid_rows.append({
                    'doctype':      'Sale Invoice Items',
                    'domain':       domain,
                    'subscription': plan_name,
                    'description':  col(row, 'description'),
                    'order_name':   col(row, 'order_name'),
                    'start_date':   parse_date(col(row, 'start_date')),
                    'end_date':     parse_date(col(row, 'end_date')),
                    'quantity':     int(float(qty_raw) if qty_raw else 0),
                    'licenses':     int(float(lic_raw)  if lic_raw  else 0),
                    'po_number':    col(row, 'po_number'),
                    'amount':       float(amt_raw if amt_raw else 0),
                })

            if not valid_rows:
                failed.append({
                    'invoice_number':        inv_num,
                    'missing_domains':       missing_domains,
                    'missing_subscriptions': missing_subscriptions,
                    'error': 'No valid rows found — all skipped due to missing domain/subscription.'
                })
                all_miss_d.extend(missing_domains)
                all_miss_s.extend(missing_subscriptions)
                continue

            try:
                doc = frappe.get_doc({
                    'doctype':        'Sale Invoice',
                    'bill_to':        col(first, 'bill_to'),
                    'invoice_number': inv_num,
                    'invoice_date':   parse_date(col(first, 'invoice_date')),
                    'due_date':       parse_date(col(first, 'due_date')),
                    'currency':       col(first, 'currency') or 'INR',
                    'gst_rate':       18,
                    'domain_details': valid_rows
                })
                doc.insert(ignore_permissions=True)
                frappe.db.commit()

                # PDF auto-match by invoice number
                pdf_attached = False
                pdf_filename = find_matching_pdf(inv_num, pdf_files)
                if pdf_filename:
                    pdf_attached = attach_pdf_to_doc(
                        doc.name, pdf_filename, pdf_files[pdf_filename]
                    )

                created.append({
                    'doc':                   doc.name,
                    'invoice_number':        inv_num,
                    'total_rows':            len(valid_rows),
                    'grand_total':           doc.grand_total,
                    'pdf_attached':          pdf_attached,
                    'pdf_filename':          pdf_filename or '',
                    'missing_domains':       missing_domains,
                    'missing_subscriptions': missing_subscriptions,
                })
                all_miss_d.extend(missing_domains)
                all_miss_s.extend(missing_subscriptions)

            except Exception as e:
                frappe.log_error(frappe.get_traceback(), 'Sale Invoice Import Error')
                failed.append({'invoice_number': inv_num, 'error': str(e)})

        return {
            'success':               True,
            'created':               created,
            'duplicates':            duplicates,
            'failed':                failed,
            'missing_domains':       list(set(all_miss_d)),
            'missing_subscriptions': list(set(all_miss_s)),
        }

    except Exception as e:
        frappe.log_error(frappe.get_traceback(), 'Sale Invoice Import Error')
        return {'success': False, 'error': str(e)}






